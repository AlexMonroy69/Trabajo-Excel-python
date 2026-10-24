from pathlib import Path
import sys

from openpyxl import load_workbook
from tabulate import tabulate


EXCEL_FILE = Path(__file__).with_name("personas.xlsx")
HEADERS = ["#", "Id", "Name", "Company", "Email", "MAC Address"]


def load_people(file_path=EXCEL_FILE):
    if not file_path.is_file():
        raise FileNotFoundError(f"No se encontró el archivo: {file_path.name}")

    workbook = load_workbook(file_path, read_only=True, data_only=True)
    worksheet = workbook.active
    people = []

    for row_number, cells in enumerate(
        worksheet.iter_rows(min_row=2, values_only=True), start=2
    ):
        if all(value is None or str(value).strip() == "" for value in cells):
            continue
        if any(value is None or str(value).strip() == "" for value in cells):
            workbook.close()
            raise ValueError(f"Hay campos vacíos en la fila {row_number}.")
        people.append([len(people) + 1, *cells])

    workbook.close()
    return people


def show_table(rows, headers=HEADERS):
    if not rows:
        print("Sin resultados.")
        return
    print(tabulate(rows, headers=headers, tablefmt="fancy_grid"))


def show_all(people):
    show_table(people)


def search_by_id(people, person_id):
    if not person_id.strip():
        print("El ID no puede estar vacío.")
        return
    matches = [person for person in people if str(person[1]).lower() == person_id.strip().lower()]
    show_table(matches)


def filter_company(people, word):
    search_word = word.strip() or "Group"
    matches = [person for person in people if search_word.lower() in str(person[3]).lower()]
    show_table(matches)


def count_records(people):
    print(f"Cantidad de registros: {len(people)}")


def show_names_and_emails(people):
    show_table([[person[2], person[4]] for person in people], ["Name", "Email"])


def run_menu(people):
    while True:
        print(
            "\n1. Ver todos\n"
            "2. Buscar por ID\n"
            "3. Filtrar empresa\n"
            "4. Contar registros\n"
            "5. Mostrar Name y Email\n"
            "6. Salir"
        )
        option = input("Elige una opción: ").strip()

        if option == "1":
            show_all(people)
        elif option == "2":
            search_by_id(people, input("Escribe el ID: "))
        elif option == "3":
            filter_company(people, input("Palabra de empresa (Enter = Group): "))
        elif option == "4":
            count_records(people)
        elif option == "5":
            show_names_and_emails(people)
        elif option == "6":
            print("Programa finalizado.")
            break
        else:
            print("Opción inválida. Elige un número del 1 al 6.")


def main():
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    try:
        people = load_people()
    except (FileNotFoundError, ValueError) as error:
        print(f"Error: {error}")
        return

    if not people:
        print("El Excel no contiene registros.")
        return
    run_menu(people)


if __name__ == "__main__":
    main()
